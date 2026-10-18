"""Обновлённый датасет после спринта — для проверки и демонстрации.

Берёт датасет в формате хакатона, строит по нему план и «проводит»
ближайший спринт: задачи получают часы и статусы по плану спринта,
команды — строку производительности в Team_history. Задачи из --delay
в спринте не двигаются — так моделируется отклонение. Результат — тот же
однолистовой Excel, который загружается в ОРБИТУ после спринта.

    python orbit/scripts/sprint_dataset.py --sprint 1 --out var/dataset_s1.xlsx
    python orbit/scripts/sprint_dataset.py --sprint 1 --delay DB-202 --out var/dataset_s1.xlsx
    python orbit/scripts/sprint_dataset.py --base var/dataset_s1.xlsx --sprint 2 --out var/dataset_s2.xlsx

Для второго и следующих спринтов основой берётся датасет предыдущего:
план по нему начинается с его среза, и первый спринт этого плана —
следующий спринт квартала.

Формулы исходного листа заменяются их значениями: так файл читается
одинаково в любой программе.
"""

import argparse
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT.parent))

from planner.engine import default_config  # noqa: E402
from planner.importer import import_xlsx  # noqa: E402
from planner.present import build_plan  # noqa: E402
from src.dataset import DATASET_PATH, SHEET_NAME  # noqa: E402
from src.normalizer import role_key  # noqa: E402

DONE_RESULT = "планируемый результат работ достигнут"


def find_row(ws, prefix, start=1):
    for row in range(start, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str) and value.strip().startswith(prefix):
            return row
    raise ValueError(f"Не найден блок {prefix}")


def simulate(base, sprint, delay=(), start_date=None, partial=None):
    """Итоги ближайшего спринта по плану: статусы, накопленные часы, SP команд.

    delay — задачи, которые в спринте не двигаются; partial — задачи,
    сделанные на долю плана спринта (0.5 — половина часов). Такие задачи
    в спринте не закрываются.

    Производительность команды за спринт = её средняя фактическая
    производительность × доля выполненного плана спринта. Команда,
    сделавшая свой план спринта целиком, работает в обычном темпе и не
    меняет среднюю, а отставание снижает её пропорционально. Писать сами
    SP плана нельзя: ёмкость плана — 80 % средней, и команда, работающая
    ровно по плану, с каждым спринтом «теряла» бы по 20 %.
    """
    data = import_xlsx(base)
    config = default_config()
    plan = build_plan(data, config)
    partial = partial or {}
    statuses, spent = {}, defaultdict(dict)
    planned, done = defaultdict(float), defaultdict(float)

    for task_id, entry in plan["entries"].items():
        task = plan["tasks"][task_id]
        if task["status"] == "Done":
            continue
        fragments = [f for f in entry["fragments"] if f["sprint"] == 1]
        if not fragments:
            continue
        for fragment in fragments:
            planned[task["team"]] += fragment["sp"]
        share = 0.0 if task_id in delay else float(partial.get(task_id, 1.0))
        if share <= 0:
            continue
        for fragment in fragments:
            done[task["team"]] += fragment["sp"] * share
            for row in fragment["assignments"]:
                spent[task_id][row["role"]] = spent[task_id].get(row["role"], 0) + row["hours"] * share
        statuses[task_id] = "Done" if entry["end"] == 1 and share >= 1 else "InProgress"

    start = date.fromisoformat(start_date or config["start_date"])
    snapshot_date = start + timedelta(days=14 * sprint - 1)
    teams = sorted({t["team"] for t in plan["tasks"].values()} | {h["team"] for h in data["history"]})
    history = defaultdict(list)
    for row in data["history"]:
        history[row["team"]].append(row["achieved"])
    velocity = {}
    for team in teams:
        average = sum(history[team]) / len(history[team]) if history[team] else 0.0
        share = done[team] / planned[team] if planned[team] > 1e-9 else 1.0
        velocity[team] = dict(achieved=round(average * share, 1), planned=round(planned[team], 1))
    return data, statuses, spent, velocity, snapshot_date


def write(base, out, data, statuses, spent, velocity, snapshot_date):
    book = load_workbook(base, data_only=True)
    ws = book[SHEET_NAME] if SHEET_NAME in book.sheetnames else book.active

    header = find_row(ws, "Номер инициативы")
    columns = {ws.cell(header, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(header, c).value}
    task_rows = {ws.cell(r, columns["task_id"]).value: r for r in range(header + 1, find_row(ws, "Spent_time_roles"))
                 if ws.cell(r, columns["task_id"]).value}

    for task_id, status in statuses.items():
        row = task_rows[task_id]
        ws.cell(row, columns["status"]).value = status
        if status == "Done":
            ws.cell(row, columns["actual_end"]).value = snapshot_date
            ws.cell(row, columns["итоговый результат"]).value = DONE_RESULT

    spent_header = find_row(ws, "task_id", find_row(ws, "Spent_time_roles"))
    roles = {role_key(ws.cell(spent_header, c).value): c for c in range(3, ws.max_column + 1) if ws.cell(spent_header, c).value}
    spent_rows = {}
    row = spent_header + 1
    while ws.cell(row, 1).value:
        spent_rows[ws.cell(row, 1).value] = row
        row += 1

    tasks = {t["id"]: t for t in data["tasks"]}
    for task_id, hours in sorted(spent.items()):
        if task_id not in spent_rows:
            ws.insert_rows(row)
            ws.cell(row, 1).value = task_id
            ws.cell(row, 2).value = tasks[task_id]["original_hh"]
            for column in roles.values():
                ws.cell(row, column).value = 0
            spent_rows[task_id] = row
            row += 1
        target = spent_rows[task_id]
        for role, value in hours.items():
            column = roles[role_key(role)]
            ws.cell(target, column).value = round((ws.cell(target, column).value or 0) + value, 3)

    history = find_row(ws, "snapshot_date", find_row(ws, "Team_history"))
    last = history + 1
    while ws.cell(last, 1).value:
        last += 1
    for team, row in velocity.items():
        ws.insert_rows(last)
        ws.cell(last, 1).value = snapshot_date
        ws.cell(last, 2).value = team
        ws.cell(last, 3).value = row["achieved"]
        ws.cell(last, 4).value = row["planned"]
        last += 1

    Path(out).parent.mkdir(parents=True, exist_ok=True)
    book.save(out)


def main():
    parser = argparse.ArgumentParser(description="Обновлённый датасет после спринта по плану")
    parser.add_argument("--base", default=str(DATASET_PATH), help="Датасет до спринта")
    parser.add_argument("--sprint", type=int, required=True, help="Номер спринта квартала")
    parser.add_argument("--delay", nargs="*", default=[], help="Задачи, которые в спринте не двигаются")
    parser.add_argument("--partial", nargs="*", default=[], metavar="ЗАДАЧА=ДОЛЯ",
                        help="Задачи, сделанные на долю плана спринта, например DB-202=0.5")
    parser.add_argument("--out", required=True, help="Куда сохранить xlsx")
    args = parser.parse_args()
    partial = {}
    for item in args.partial:
        task_id, _, share = item.partition("=")
        partial[task_id] = float(share)
    data, statuses, spent, velocity, snapshot_date = simulate(args.base, args.sprint, set(args.delay), partial=partial)
    write(args.base, args.out, data, statuses, spent, velocity, snapshot_date)
    done = sorted(t for t, s in statuses.items() if s == "Done")
    print(f"Спринт {args.sprint}: закрыто {len(done)} задач, в работе {len(statuses) - len(done)}, "
          f"отложено {len(args.delay)}; производительность команд: "
          f"{ {team: row['achieved'] for team, row in velocity.items()} }")
    print(f"Сохранено: {args.out}")


if __name__ == "__main__":
    main()
